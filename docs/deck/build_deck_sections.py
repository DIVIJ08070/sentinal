#!/usr/bin/env python3
"""SENTINEL deck — sectioned (tech stack / what it does / problems / results /
achieved / future). Gujarat CCTV Hackathon 2026. Outputs the submission PPTX."""
import os
from pptx import Presentation
from pptx.util import Inches as IN, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

_HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(_HERE, "..", "..", "deliverables", "SENTINEL_Presentation.pptx")

NAVY = RGBColor(0x0B, 0x12, 0x20); CARD = RGBColor(0x13, 0x1C, 0x31)
CARD2 = RGBColor(0x18, 0x23, 0x38); LINE = RGBColor(0x2A, 0x38, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); BODY = RGBColor(0xE6, 0xED, 0xF7)
MUT = RGBColor(0x8F, 0xA1, 0xBC); AMBER = RGBColor(0xF5, 0xA6, 0x23)
GREEN = RGBColor(0x34, 0xD3, 0x99); RED = RGBColor(0xF8, 0x71, 0x71)
FONT = "Arial"; MONO = "Consolas"
W, H = 13.333, 7.5

prs = Presentation(); prs.slide_width = IN(W); prs.slide_height = IN(H)
BLANK = prs.slide_layouts[6]
_pageno = [0]


def bg(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY


def rect(s, x, y, w, h, fill=CARD, line=None, radius=0.06, oval=False):
    shp = MSO_SHAPE.OVAL if oval else MSO_SHAPE.ROUNDED_RECTANGLE
    sp = s.shapes.add_shape(shp, IN(x), IN(y), IN(w), IN(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    if not oval:
        try: sp.adjustments[0] = radius
        except Exception: pass
    return sp


def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tf = s.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h)).text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, runs, size=12, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         before=0, after=0, first=False, line=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if before: p.space_before = Pt(before)
    if after: p.space_after = Pt(after)
    if line: p.line_spacing = line
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, o in runs:
        r = p.add_run(); r.text = text
        r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", bold)
        r.font.italic = o.get("italic", False)
        r.font.color.rgb = o.get("color", color); r.font.name = o.get("font", FONT)
    return p


def slide(kicker, title, sub=None):
    s = prs.slides.add_slide(BLANK); bg(s)
    rect(s, 0, 0, 0.18, H, fill=AMBER, radius=0)
    para(box(s, 0.55, 0.42, 12.4, 0.32), kicker.upper(), size=12.5, color=AMBER, bold=True, first=True)
    para(box(s, 0.55, 0.74, 12.4, 0.64), title, size=26, color=WHITE, bold=True, first=True)
    if sub:
        para(box(s, 0.55, 1.4, 12.4, 0.34), [(sub, {"italic": True})], size=13.5, color=MUT, first=True)
    _pageno[0] += 1
    n = _pageno[0]
    para(box(s, 0.55, 7.12, 10.0, 0.3),
         "SENTINEL  ·  Divij Patel · Meet Patel · Krish Patel  ·  Gujarat Police CCTV Hackathon 2026",
         size=9, color=MUT, first=True)
    para(box(s, 12.2, 7.12, 0.9, 0.3), f"{n:02d}", size=9, color=MUT, align=PP_ALIGN.RIGHT, first=True)
    return s


def chip(s, x, y, w, text, color=AMBER, fill=CARD2, h=0.34, size=10.5, bold=True):
    c = rect(s, x, y, w, h, fill=fill, radius=0.5)
    tf = c.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = IN(0.08); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT; r.font.bold = bold
    return c


def card(s, x, y, w, h, title, lines, title_color=AMBER, fill=CARD):
    rect(s, x, y, w, h, fill=fill, line=LINE)
    tf = box(s, x + 0.22, y + 0.16, w - 0.44, h - 0.32)
    para(tf, title, size=13.5, color=title_color, bold=True, first=True)
    for ln in lines:
        if isinstance(ln, tuple):
            para(tf, [("• ", {"color": MUT}), ln], size=11, color=BODY, before=3, line=1.05)
        else:
            para(tf, [("• ", {"color": MUT}), (ln, {})], size=11, color=BODY, before=3, line=1.05)


def stat(s, x, y, w, big, label, color=GREEN):
    rect(s, x, y, w, 1.15, fill=CARD, line=LINE)
    para(box(s, x, y + 0.14, w, 0.6, MSO_ANCHOR.MIDDLE),
         big, size=25, color=color, bold=True, align=PP_ALIGN.CENTER, first=True)
    para(box(s, x + 0.1, y + 0.72, w - 0.2, 0.36),
         label, size=9.5, color=MUT, align=PP_ALIGN.CENTER, first=True)


# ===================================================== 1 — TITLE
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, W, 0.14, fill=AMBER, radius=0)
para(box(s, 0.9, 1.5, 11.5, 0.4), "GUJARAT POLICE CCTV HACKATHON 2026  ·  CATEGORY 1",
     size=13, color=AMBER, bold=True, first=True)
para(box(s, 0.9, 2.05, 11.5, 1.1), "SENTINEL", size=58, color=WHITE, bold=True, first=True)
para(box(s, 0.9, 3.15, 11.5, 0.6),
     [("Plate to court in sixty seconds — on the cameras Gujarat already owns.", {"italic": True})],
     size=19, color=AMBER, bold=True, first=True)
para(box(s, 0.9, 3.95, 11.6, 0.8),
     "A unified CCTV registry, GIS, live viewing, CPU-only ANPR, watchlist alerts, "
     "route reconstruction and court-ready evidence — integrating heterogeneous feeds "
     "from 26 departments without replacing a single camera.", size=13.5, color=BODY, line=1.15, first=True)
chip(s, 0.9, 5.05, 3.0, "Hybrid Model 1 + 2/4", color=BODY)
chip(s, 4.05, 5.05, 2.85, "CPU-only · open-source", color=BODY)
chip(s, 7.0, 5.05, 2.95, "SHA-256 evidence dossier", color=BODY)
para(box(s, 0.9, 6.1, 11.5, 0.4),
     [("TEAM:  ", {"bold": True, "color": MUT}),
      ("Divij Patel · Meet Patel · Krish Patel", {"color": BODY}),
      ("   ·   vatsunp11@gmail.com", {"color": MUT})], size=13, first=True)

# ===================================================== 2 — WHAT IT DOES
s = slide("What it does", "Type a registration number, get a court-ready answer",
          "One platform: register · view · detect · correlate · trace · prove.")
beats = [
    ("1 · REGISTER & VIEW", "Catalogue-driven registry + GIS map of every department's cameras — metadata, health, CSV import/export, and a gap analysis that flags coverage blind-spots. Live viewing of heterogeneous feeds; nothing hard-coded, nothing replaced."),
    ("2 · DETECT (ANPR)", "CPU-only vehicle detection + plate reading on the live feed, every read PTS-timestamped with a displayed confidence."),
    ("3 · CORRELATE & ALERT", "Each read checked against the watchlist; a match fires a real-time alert in under a second with the camera's own frame."),
    ("4 · TRACE THE ROUTE", "Search a plate → the timestamped, camera-by-camera route on the GIS map, with a physics filter rejecting impossible hops."),
    ("5 · PROVE IT", "One click exports a SHA-256 hash-chained chain-of-custody dossier — the document an officer attaches to a charge sheet."),
]
y = 1.95
for i, (t, d) in enumerate(beats):
    rect(s, 0.55, y, 12.25, 0.92, fill=CARD if i % 2 == 0 else CARD2, line=LINE)
    para(box(s, 0.8, y + 0.14, 3.3, 0.64, MSO_ANCHOR.MIDDLE), t, size=13, color=AMBER, bold=True, first=True)
    para(box(s, 4.2, y + 0.12, 8.4, 0.7, MSO_ANCHOR.MIDDLE), d, size=11.5, color=BODY, line=1.05, first=True)
    y += 1.0

# ===================================================== 3 — PROBLEMS SOLVED
s = slide("Problems solved", "The control-room pain, and how SENTINEL answers it")
probs = [
    ("Fragmented, incompatible CCTV (26 depts)", "One catalogue-driven registry + GIS; integrate don't replace; mixed H.264/H.265 and vendors handled per-camera."),
    ("No view of where coverage is missing", "A gap analysis flags districts and areas with inadequate CCTV — with an exportable report for planning."),
    ("Days of manual footage scrubbing", "Plate in → timestamped multi-camera route out in seconds."),
    ("OCR misreads on real CCTV plates", "RapidOCR + a deterministic grammar-aware arbiter (“the king”) + confidence + TrustTrace."),
    ("Evidence that dies before court", "SHA-256 hash-chained dossier; snapshots bound by hash; operator identity on every record."),
    ("Insider tampering with records", "Append-only, hash-chained audit log — any edit breaks the chain and is caught on verify."),
    ("Scaling to 80,000 cameras / 160 Gbps", "Edge-first: video stays on departmental NVRs; ~1–3 Kbps of metadata travels upstream."),
    ("Vendor lock-in and cost", "CPU-only, all open-source; GPUs optional; Make-in-India procurement path."),
]
for i, (p, a) in enumerate(probs):
    col = i % 2; row = i // 2
    x = 0.55 + col * 6.2; y = 1.85 + row * 1.28
    rect(s, x, y, 5.95, 1.12, fill=CARD, line=LINE)
    para(box(s, x + 0.2, y + 0.14, 5.6, 0.4), [("✗  ", {"color": RED, "bold": True}), (p, {"bold": True, "color": WHITE})], size=12, first=True)
    para(box(s, x + 0.2, y + 0.54, 5.6, 0.5), [("→  ", {"color": GREEN, "bold": True}), (a, {"color": BODY})], size=10.5, line=1.03, first=True)

# ===================================================== 4 — TECH STACK
s = slide("Tech stack", "All open-source, CPU-only, no vendor lock, no GPU dependency")
card(s, 0.55, 1.85, 3.95, 2.35, "Frontend", [
    "React 18 + Vite", "Leaflet GIS map", "HLS video (hls.js)", "Real-time alerts over WebSocket"])
card(s, 4.7, 1.85, 3.95, 2.35, "Backend", [
    "Python 3.12 · FastAPI", "SQLAlchemy ORM", "WebSockets (live alerts)", "fpdf2 (evidence dossier)"])
card(s, 8.85, 1.85, 3.95, 2.35, "Data & registry", [
    "PostgreSQL + PostGIS (prod)", "Camera registry + gap analysis", "CSV import/export · dept scope", "Hash-chained audit · SHA-256 dossier"])
card(s, 0.55, 4.35, 3.95, 2.45, "AI / ANPR (CPU)", [
    "YOLOv8n vehicle detection", "Plate localizer (ONNX)", "RapidOCR + Indian syntax gate", "“The king” — deterministic", "   grammar-aware arbiter"], title_color=GREEN)
card(s, 4.7, 4.35, 3.95, 2.45, "Streaming", [
    "RTSP forced over TCP", "ffmpeg HLS relay", "PTS-anchored timestamps", "Reconnect w/ backoff 2→30 s"])
card(s, 8.85, 4.35, 3.95, 2.45, "Infra & scale", [
    "Docker + Docker Compose", "Kubernetes manifests + HPA", "Redpanda (Kafka API) event bus", "SeaweedFS (S3 object storage)"], title_color=AMBER)

# ===================================================== 5 — RESULTS
s = slide("Results", "Measured, not asserted — on live government feeds and a running PoC")
stat(s, 0.55, 1.85, 2.9, "30 / 30", "gov cameras onboarded live", GREEN)
stat(s, 3.6, 1.85, 2.9, "1 → 11 / 19", "ANPR exact reads (RapidOCR)", AMBER)
stat(s, 6.65, 1.85, 2.9, "6 / 6", "fail cases the king recovered", GREEN)
stat(s, 9.7, 1.85, 3.1, "< 1 s", "watchlist alert, real-time", GREEN)
stat(s, 0.55, 3.2, 2.9, "16 ms", "Postgres replication lag (DR)", AMBER)
stat(s, 3.6, 3.2, 2.9, "15,600/s", "Kafka events · 1.7 ms p50", AMBER)
stat(s, 6.65, 3.2, 2.9, "~40×", "statewide event headroom", GREEN)
stat(s, 9.7, 3.2, 3.1, "345/s", "SeaweedFS object writes", AMBER)
card(s, 0.55, 4.65, 6.0, 2.1, "Evidence & trust (built)", [
    "Tamper-evident audit — hash chain, verifiable by recomputation",
    "SHA-256 chain-of-custody dossier, snapshots bound by hash",
    "TrustTrace: confidence, cross-read agreement, cloned-plate flag",
    "Physics filter rejects impossible hops with a plain reason"], title_color=GREEN)
card(s, 6.75, 4.65, 6.05, 2.1, "Registry & scale (computed)", [
    "Gap analysis: 30 cameras mapped across 7 of 33 districts",
    "Edge-first: 160 Gbps video → ~80–240 Mbps events statewide",
    "150 TB of snapshots vs ~52 PB/month of video (rejected)",
    "18 backend tests passing; full Docker stack on PostgreSQL"], title_color=AMBER)

# ===================================================== 6 — WHAT WE ACHIEVED
s = slide("What we achieved", "A working, honest, production-shaped platform")
card(s, 0.55, 1.85, 6.0, 2.3, "Complete & verified", [
    "Every mandatory module built and running on live feeds",
    "Model 1 registry: metadata · gap analysis · CSV · dept scope",
    "ANPR · watchlist · alerts · route + physics · dossier",
    "18 backend tests pass; clean-clone quickstart"], title_color=GREEN)
card(s, 6.75, 1.85, 6.05, 2.3, "Evidence-first, by design", [
    "Never asserts a read as fact — confidence on everything",
    "Confusion-tolerant matcher recovers OCR misreads",
    "TrustTrace explains every result for human verification",
    "Tamper-evident audit + hash-chained dossier"], title_color=AMBER)
card(s, 0.55, 4.3, 6.0, 2.45, "Engineering we added", [
    "“The king” — deterministic, open-source plate reconciliation",
    "RapidOCR + Indian-plate syntax gate (10-char capable)",
    "SQLite → PostgreSQL; full Docker one-command deploy",
    "Production-component PoC: replica + Kafka + object store"], title_color=GREEN)
card(s, 6.75, 4.3, 6.05, 2.45, "Honest & affordable", [
    "CPU-only core — GPUs optional, Make-in-India friendly",
    "All open-source; no vendor lock, no LLM dependency",
    "~₹15–20 cr statewide capex; zero new cameras",
    "Built vs designed stated plainly throughout"], title_color=AMBER)

# ===================================================== 7 — FUTURE IMPLEMENTATION
s = slide("Future implementation", "The roadmap — proven by the PoC, not just drawn")
fut = [
    ("Sharper OCR", "Perspective rectification of angled plates, a multi-engine ensemble, and optional super-resolution — all feeding the king, all CPU/open-source."),
    ("Event bus in the live path", "Move detections onto the Kafka (Redpanda) bus that the PoC already proved at 40× statewide headroom — buffering, fan-out, replay."),
    ("DR & object storage", "PostgreSQL streaming replica + a DR site (16 ms lag demonstrated); snapshots to S3 object storage (SeaweedFS)."),
    ("Authorised DB integration", "Read-only, audit-logged adapters for VAHAN / SARTHI / eGujCop — owner, licence and FIR-linked watchlist auto-population."),
    ("Edge / regional tiering", "Inference at ~30 regional PoPs on GSWAN; edge boxes only for thin rural links. Video never crosses the backbone."),
    ("Phased statewide rollout", "500 → 5,000 → 25,000 → 80,000 cameras, with new analytics (crowd, anomaly) at software-only marginal cost."),
]
for i, (t, d) in enumerate(fut):
    col = i % 2; row = i // 2
    x = 0.55 + col * 6.2; y = 1.95 + row * 1.6
    rect(s, x, y, 5.95, 1.42, fill=CARD, line=LINE)
    rect(s, x, y, 0.1, 1.42, fill=AMBER, radius=0)
    para(box(s, x + 0.28, y + 0.16, 5.5, 0.4), t, size=13.5, color=AMBER, bold=True, first=True)
    para(box(s, x + 0.28, y + 0.58, 5.5, 0.78), d, size=11, color=BODY, line=1.07, first=True)

# ===================================================== 8 — CLOSE
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 3.0, W, 0.03, fill=LINE, radius=0)
para(box(s, 0.9, 2.1, 11.5, 0.9, MSO_ANCHOR.BOTTOM),
     [("Plate to court in sixty seconds —", {"color": WHITE, "bold": True})],
     size=30, align=PP_ALIGN.CENTER, first=True)
para(box(s, 0.9, 3.15, 11.5, 0.7),
     [("on the cameras Gujarat already owns.", {"color": AMBER, "bold": True})],
     size=30, align=PP_ALIGN.CENTER, first=True)
para(box(s, 0.9, 4.3, 11.5, 0.5),
     "Zero new cameras.  CPU-only.  Open-source.  Evidence-first.",
     size=15, color=MUT, align=PP_ALIGN.CENTER, first=True)
para(box(s, 0.9, 6.4, 11.5, 0.4),
     [("Divij Patel · Meet Patel · Krish Patel", {"color": BODY, "bold": True}),
      ("   ·   Gujarat Police CCTV Hackathon 2026 · Category 1", {"color": MUT})],
     size=13, align=PP_ALIGN.CENTER, first=True)

prs.core_properties.author = "Divij Patel, Meet Patel, Krish Patel"
prs.core_properties.title = "SENTINEL — Gujarat Police CCTV Hackathon 2026"
prs.core_properties.subject = "Gujarat Police CCTV Hackathon 2026 — Category 1, Team"
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides._sldIdLst))
